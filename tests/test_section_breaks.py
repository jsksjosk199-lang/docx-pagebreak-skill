"""Synthetic regression tests for the public section-break helper."""
# Copyright 2026 jsksjosk199-lang
# SPDX-License-Identifier: Apache-2.0
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


SCRIPT = Path(__file__).resolve().parents[1] / "docx-pagebreak" / "scripts" / "section_breaks.py"


def make_report(path, *, boundary_symbol=False, duplicate=False, inherited=False):
    doc = Document()
    first = doc.add_table(rows=2, cols=1)
    first.cell(0, 0).text = "本月完成工作"
    first.cell(1, 0).text = "示例工作内容"
    boundary = doc.add_paragraph()
    if boundary_symbol:
        symbol = OxmlElement("w:sym")
        symbol.set(qn("w:font"), "Symbol")
        symbol.set(qn("w:char"), "F0A7")
        boundary.add_run()._r.append(symbol)
    plan = doc.add_table(rows=2, cols=1)
    title = plan.cell(0, 0).paragraphs[0]
    title.add_run("下月工作计划").add_break(WD_BREAK.PAGE)
    plan.cell(1, 0).text = "示例计划内容"
    plan._element.tblPr.append(OxmlElement("w:tblpPr"))
    if inherited:
        style = doc.styles.add_style("SyntheticBreakTitle", WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.page_break_before = True
        title.style = style
    doc.add_paragraph()
    review = doc.add_table(rows=2, cols=1)
    review.cell(0, 0).text = "问题复盘与总结"
    review.cell(1, 0).text = "示例复盘内容"
    if duplicate:
        extra = doc.add_table(rows=1, cols=1)
        extra.cell(0, 0).text = "下月工作计划"
    doc.save(path)


def document_xml(path):
    with ZipFile(path) as docx:
        return docx.read("word/document.xml")


def canonical_xml(path):
    return etree.tostring(etree.fromstring(document_xml(path)), method="c14n")


class SectionBreakTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        self.output = self.root / "result.docx"

    def run_script(self, *args):
        env = dict(os.environ, PYTHONIOENCODING="utf-8")
        return subprocess.run(
            [sys.executable, str(SCRIPT), *map(str, args)],
            text=True, encoding="utf-8", capture_output=True, env=env,
            check=False,
        )

    def inspect(self, path):
        result = self.run_script("inspect", path)
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def test_both_boundaries_move_heading_break_and_inline_floating_table(self):
        make_report(self.source, inherited=True)
        before = self.inspect(self.source)
        self.assertEqual(before["embedded_header_page_breaks"], 1)
        self.assertTrue(before["sections"][1]["floating"])
        result = self.run_script("repair", self.source, "--out", self.output,
                                 "--boundaries", "both", "--inline-floating")
        self.assertEqual(result.returncode, 0, result.stderr)
        changes = json.loads(result.stdout)
        self.assertEqual(changes["embedded_breaks_moved"], 1)
        self.assertEqual(changes["heading_page_break_before_cleared"], 1)
        after = self.inspect(self.output)
        self.assertEqual(after["embedded_header_page_breaks"], 0)
        self.assertFalse(after["sections"][1]["floating"])
        self.assertEqual([b["manual_page_breaks"] for b in after["boundaries"]], [1, 1])
        texts = [cell.text for table in Document(self.output).tables for row in table.rows for cell in row.cells]
        self.assertTrue({"示例工作内容", "示例计划内容", "示例复盘内容"}.issubset(set(texts)))

    def test_only_requested_boundary_is_changed(self):
        make_report(self.source)
        before_review = Document(self.source).tables[2]._element.xml
        result = self.run_script("repair", self.source, "--out", self.output,
                                 "--boundaries", "month-plan")
        self.assertEqual(result.returncode, 0, result.stderr)
        after = self.inspect(self.output)
        self.assertEqual([b["manual_page_breaks"] for b in after["boundaries"]], [1, 0])
        self.assertEqual(Document(self.output).tables[2]._element.xml, before_review)

    def test_repair_is_idempotent(self):
        make_report(self.source)
        second = self.root / "second.docx"
        first_run = self.run_script("repair", self.source, "--out", self.output,
                                    "--boundaries", "both", "--wps-center", "--inline-floating")
        self.assertEqual(first_run.returncode, 0, first_run.stderr)
        second_run = self.run_script("repair", self.output, "--out", second,
                                     "--boundaries", "both", "--wps-center", "--inline-floating")
        self.assertEqual(second_run.returncode, 0, second_run.stderr)
        self.assertEqual(canonical_xml(self.output), canonical_xml(second))

    def test_symbol_at_boundary_is_reported_and_not_removed(self):
        make_report(self.source, boundary_symbol=True)
        self.assertFalse(self.inspect(self.source)["boundaries"][0]["supported"])
        result = self.run_script("repair", self.source, "--out", self.output,
                                 "--boundaries", "both")
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())
        self.assertIn(b"w:sym", document_xml(self.source))

    def test_ambiguous_headings_and_existing_output_are_rejected(self):
        make_report(self.source, duplicate=True)
        result = self.run_script("repair", self.source, "--out", self.output,
                                 "--boundaries", "both")
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())
        make_report(self.source)
        self.output.write_text("untouched", encoding="utf-8")
        result = self.run_script("repair", self.source, "--out", self.output,
                                 "--boundaries", "both")
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.output.read_text(encoding="utf-8"), "untouched")


if __name__ == "__main__":
    unittest.main()
